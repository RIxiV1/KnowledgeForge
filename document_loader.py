import os
import re
import json
import fitz  # PyMuPDF
import pandas as pd
from langchain_core.documents import Document
from docx import Document as DocxDocument
from pptx import Presentation


def _clean(text):
    """Drop only genuinely undecodable glyphs and collapse runs of whitespace.

    Real punctuation (em/en dashes, bullets, arrows, curly quotes) is kept —
    it's valid content; only the Windows console fails to render it.
    """
    text = text.replace("�", " ")  # replacement char = genuinely undecodable glyph
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

CSV_BATCH_SIZE = 100
def create_metadata(filename, file_path, file_type):
    return { "source": filename, "file_path": file_path, "file_type": file_type }

# PDF And TXT

def load_file(file_path):
    extension = os.path.splitext(file_path)[1].lower()
    filename = os.path.basename(file_path)
    docs = []
    if extension == ".pdf":
        # PyMuPDF reconstructs word spacing from glyph positions far better than
        # the default loader (which jams words together on styled/web-made PDFs).
        pdf = fitz.open(file_path)
        for i, page in enumerate(pdf):
            text = _clean(page.get_text("text"))
            if not text:
                continue
            docs.append(
                Document(
                    page_content=text,
                    metadata={**create_metadata(filename, file_path, "pdf"), "page": i + 1},
                )
            )
        pdf.close()
    elif extension == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            text = _clean(f.read())
        # NOTE: no dataset_id here — that tag marks a file as a tabular dataset
        # for the analytics path. A .txt is prose, so tagging it made questions
        # like "how many entries" route into the spreadsheet analyzer, which
        # then answered with "Unsupported file format for analytics".
        docs = [
            Document(page_content=text, metadata=create_metadata(filename, file_path, "txt"))
        ]
# CSV

    elif extension == ".csv":
        df = pd.read_csv(file_path, low_memory=False)
        total_rows = len(df)

        for start in range(0, total_rows, CSV_BATCH_SIZE):
            batch_df = df.iloc[start:start + CSV_BATCH_SIZE]
            docs.append(
                Document(
                    page_content=batch_df.to_string(index=False),
                    metadata={
                        **create_metadata(filename, file_path, "csv"),
                        "dataset_id": file_path,
                        "batch_start": start,
                        "batch_end": min(start + CSV_BATCH_SIZE, total_rows)
                    }
                )
            )

# Excel Format (xlsx, xls)
    elif extension in [".xlsx", ".xls"]:
        df = pd.read_excel(file_path)
        total_rows = len(df)
        for start in range(0, total_rows, CSV_BATCH_SIZE):
            batch_df = df.iloc[start:start + CSV_BATCH_SIZE]
            docs.append(
                Document( page_content=batch_df.to_string(index=False), metadata={  **create_metadata(filename, file_path, "xlsx"),  "dataset_id": file_path,  "batch_start": start,  "batch_end": min(start + CSV_BATCH_SIZE, total_rows)} ) )

# DOCX

    elif extension == ".docx":
        # Split on heading paragraphs so each section is its own document. This
        # keeps unrelated topics out of the same chunk (better retrieval) and
        # lets citations point to the section a fact came from. Falls back to a
        # single blob for documents with no heading styles.
        doc = DocxDocument(file_path)
        sections, current_head, current = [], None, []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if not txt:
                continue
            style = (p.style.name if p.style else "") or ""
            if style.startswith("Heading") or style == "Title":
                if current:
                    sections.append((current_head, current))
                current_head, current = txt, [txt]
            else:
                current.append(txt)
        if current:
            sections.append((current_head, current))

        if not sections:
            text = _clean("\n".join(p.text for p in doc.paragraphs))
            docs = [Document(page_content=text, metadata=create_metadata(filename, file_path, "docx"))]
        else:
            for head, paras in sections:
                text = _clean("\n".join(paras))
                if not text:
                    continue
                meta = create_metadata(filename, file_path, "docx")
                if head:
                    meta["section"] = head[:120]
                docs.append(Document(page_content=text, metadata=meta))

# PPTX

    elif extension == ".pptx":
        # One document per slide, tagged with its slide number (as "page") so
        # citations can say which slide a fact came from.
        presentation = Presentation(file_path)
        for i, slide in enumerate(presentation.slides, 1):
            parts = [shape.text for shape in slide.shapes
                     if hasattr(shape, "text") and shape.text.strip()]
            text = _clean("\n".join(parts))
            if not text:
                continue
            docs.append(
                Document(
                    page_content=text,
                    metadata={**create_metadata(filename, file_path, "pptx"), "page": i},
                )
            )

# JSON

    elif extension == ".json":
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        docs = [
            Document( page_content=json.dumps(data, indent=2), metadata=create_metadata(filename, file_path, "json") )
        ]
    else:
        raise ValueError(f"Unsupported file type: {extension}")
    return docs